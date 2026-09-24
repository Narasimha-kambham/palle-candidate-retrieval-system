import pymupdf
from docx import Document

def extract_text_from_pdf(filepath):

    text = []

    with pymupdf.open(filepath) as doc:
        for page in doc:
            text.append(page.get_text())

    return "\n".join(text)

def extract_text_from_docx(filepath):
    doc = Document(filepath)

    text = []

    for paragraph in doc.paragraphs:
        paragraph_text  = paragraph.text.strip()

        if paragraph_text:
            text.append(paragraph_text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                row_text.append(cell_text)
            if any(row_text):
                text.append("|".join(row_text))
            

    return "\n".join(text)

def extract_text_from_text(input_text):
    if not isinstance(input_text, str):
        raise TypeError("Input Text Must be a string")

    input_text = input_text.strip()

    if not input_text:
        raise ValueError("Input Text cant be empty")

    return input_text


def extract_document_text(filepath=None, raw_text=None):
    """
    Extract text from PDF, DOCX, TXT files, or raw text input.
    """
    if raw_text is not None and str(raw_text).strip():
        return extract_text_from_text(str(raw_text))

    if not filepath:
        raise ValueError("Either a file path or raw text must be provided.")

    from pathlib import Path
    path = Path(filepath)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(str(path))
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(str(path))
    elif suffix in (".txt", ".text", ".md"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return extract_text_from_text(f.read())
    else:
        raise ValueError(
            f"Unsupported document format: {suffix}. "
            "Please provide a .pdf, .docx, or .txt file."
        )
